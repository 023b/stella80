#!/usr/bin/env python3
"""Local document -> Obsidian knowledge graph agent.

Sources (pdf/docx/pptx/md and, with src_all, most other files under src_root, plus the vault's own
topic notes) are chunked, sent to a local Ollama model, and concepts are accumulated in SQLite.
The _auto/ and _raw/ folders of the vault are rendered deterministically from that state, so they
can be rebuilt at any time. Notes you wrote yourself are never edited.
"""
import argparse, fcntl, hashlib, json, os, re, sqlite3, subprocess, sys, time
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
import requests
import yaml

HERE = Path(__file__).parent
CFG_FILE = Path(os.environ.get("VAULT_AGENT_CONFIG", HERE / "config.yaml")).resolve()
CFG = yaml.safe_load(CFG_FILE.read_text())


def cfg_path(p):
    """Config paths may use ~ and are relative to the config file."""
    p = Path(os.path.expanduser(str(p)))
    return (p if p.is_absolute() else CFG_FILE.parent / p).resolve()


for k, v in {"courses": {}, "exclude_src": [], "exclude_vault_names": [], "exclude_vault_prefixes": ["00-"],
             "src_exts": [".pdf", ".docx", ".pptx", ".md"], "keep_alive": "10m", "chunk_chars": 5000,
             "merge_cosine": 0.93, "min_mentions_for_note": 2, "embed_model": "nomic-embed-text"}.items():
    CFG.setdefault(k, v)
CFG["courses"] = CFG["courses"] or {}
SRC = cfg_path(CFG["src_root"])
VAULT = cfg_path(CFG["vault_root"]) / CFG["vault_dir"]
AUTO = VAULT / "_auto"
STATE = cfg_path(CFG["state_dir"])
STATE.mkdir(parents=True, exist_ok=True)
OLL = os.environ.get("OLLAMA_URL") or CFG.get("ollama_url", "http://127.0.0.1:11434")
BAD_CHARS = re.compile(r'[\\/:*?"<>|#^\[\]]')
VP = CFG["vault_dir"] + "/"  # Obsidian vault root is vault_root, so links carry this prefix
USER_MARK = "<!-- user-notes: anything below this line is preserved -->"

RULES = yaml.safe_load(Path(os.environ.get("VAULT_AGENT_RULES") or cfg_path(CFG.get("rules", HERE / "rules.yaml"))).read_text()) or {}
RG, RT, RD = RULES["graph"], RULES["throttle"], RULES["dedupe"]


class Offline(Exception):
    pass


def log(*a):
    print(time.strftime("%H:%M:%S"), *a, flush=True)


# ---------------------------------------------------------------- state
def db():
    c = sqlite3.connect(STATE / "state.db")
    c.executescript("""
    CREATE TABLE IF NOT EXISTS docs(id INTEGER PRIMARY KEY, path TEXT UNIQUE, sha TEXT, kind TEXT,
        course TEXT, title TEXT, status TEXT, doctype TEXT, link TEXT);
    CREATE TABLE IF NOT EXISTS chunks(doc_id INT, idx INT, heading TEXT, summary TEXT, PRIMARY KEY(doc_id, idx));
    CREATE TABLE IF NOT EXISTS concepts(id INTEGER PRIMARY KEY, key TEXT UNIQUE, name TEXT, kind TEXT,
        aliases TEXT DEFAULT '[]', emb BLOB);
    CREATE TABLE IF NOT EXISTS mentions(concept_id INT, doc_id INT, idx INT, quote TEXT, definition TEXT,
        PRIMARY KEY(concept_id, doc_id, idx));
    CREATE TABLE IF NOT EXISTS chunk_cache(hash TEXT PRIMARY KEY, emb BLOB, summary TEXT, items TEXT);
    CREATE TABLE IF NOT EXISTS raw_exports(doc_id INTEGER PRIMARY KEY, sha TEXT, path TEXT);
    CREATE TABLE IF NOT EXISTS figure_pages(doc_id INT, page INT, status TEXT, PRIMARY KEY(doc_id, page));
    """)
    if "priority" not in [r[1] for r in c.execute("PRAGMA table_info(docs)")]:
        c.execute("ALTER TABLE docs ADD COLUMN priority INTEGER DEFAULT 0")   # set from the STELLA-80 admin queue: higher runs first
        c.commit()
    return c


def sha_of(p: Path):
    h = hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


# ---------------------------------------------------------------- ollama
def ensure_tunnel():
    """Check Ollama is reachable; if tunnel_cmd is set (Ollama on another machine), (re)open the tunnel first.
    Ollama can take a while to answer while it loads a large model or the machine is short of memory, so be patient."""
    for attempt in range(5):
        try:
            requests.get(OLL + "/api/tags", timeout=30).raise_for_status()
            return
        except Exception:
            if attempt == 0 and CFG.get("tunnel_cmd"):
                subprocess.run(CFG["tunnel_cmd"], shell=True, capture_output=True, timeout=30)
            time.sleep(10)
    raise Offline("Ollama unreachable")


def need_models(*keys):
    """Stop before touching any document if a model is not pulled yet, so a fresh install doesn't mark its files as failed."""
    have = {m["name"] for m in requests.get(OLL + "/api/tags", timeout=30).json().get("models", [])}
    have |= {n.removesuffix(":latest") for n in have}
    missing = [CFG[k] for k in keys if CFG[k] not in have]
    if missing:
        raise Offline("model not pulled yet: " + ", ".join(missing) + " (ollama pull <name>)")


STUDY_OLLAMA = CFG.get("study_ollama_url")   # STELLA-80's Ollama instance when it is a separate one that shares this machine's RAM
_last_free = [0.0]


def free_study_models(force=False):
    """The agent's model and STELLA-80's chat model do not fit in RAM together. When memory is short, unload the big study model
    (it reloads on the next question); small models are left alone."""
    if not STUDY_OLLAMA or STUDY_OLLAMA == OLL or (not force and time.time() - _last_free[0] < 60):
        return
    _last_free[0] = time.time()
    try:
        avail = next(int(l.split()[1]) for l in Path("/proc/meminfo").read_text().splitlines() if l.startswith("MemAvailable")) / 2**20
        if avail > 20 and not force:
            return
        for m in requests.get(STUDY_OLLAMA + "/api/ps", timeout=3).json().get("models", []):
            if m.get("size", 0) > 8e9:
                requests.post(STUDY_OLLAMA + "/api/generate", json={"model": m["name"], "keep_alive": 0}, timeout=30)
                log(f"freed study model {m['name']} to make room for the agent's model")
    except Exception:
        pass


def machine_load():
    """1-minute load average of the machine running Ollama: this one, or load_cmd's output when Ollama is elsewhere."""
    if CFG.get("load_cmd"):
        try:
            return float(subprocess.run(CFG["load_cmd"], shell=True, capture_output=True, text=True, timeout=20).stdout.split()[0])
        except Exception:
            return 0.0
    try:
        return os.getloadavg()[0]
    except OSError:
        return 0.0


def breathe(sleep_s):
    """Pause between model calls and back off while the machine is busy with something else."""
    time.sleep(sleep_s)
    waited = 0
    while machine_load() > (RT.get("max_load") or RT.get("max_stella_load") or os.cpu_count() or 8):
        if waited >= 600:
            raise Offline("the machine has been busy for 10 minutes")
        log("machine busy, waiting")
        time.sleep(60)
        waited += 60


def unload(model):
    try:
        requests.post(OLL + "/api/generate", json={"model": model, "keep_alive": 0}, timeout=60)
    except Exception:
        pass


def chat(messages, schema, model=None, images=None):
    free_study_models()
    if images:
        messages[-1]["images"] = images
    body = {"model": model or CFG["chat_model"], "messages": messages, "stream": False, "think": False,
            "format": schema, "keep_alive": CFG["keep_alive"],
            "options": {"temperature": 0.1, "num_ctx": 8192, **({"num_thread": RT["num_thread"]} if RT.get("num_thread") else {})}}
    for attempt in range(3):
        try:
            r = requests.post(OLL + "/api/chat", json=body, timeout=1800)
            r.raise_for_status()
            return json.loads(r.json()["message"]["content"])
        except (requests.ConnectionError, requests.Timeout):
            ensure_tunnel()
        except (json.JSONDecodeError, KeyError):
            pass
    return None


def embed(texts):
    r = requests.post(OLL + "/api/embed", json={"model": CFG["embed_model"], "input": texts,
                                                 "keep_alive": CFG["keep_alive"]}, timeout=300)
    r.raise_for_status()
    v = np.array(r.json()["embeddings"], dtype=np.float32)
    return v / np.maximum(np.linalg.norm(v, axis=1, keepdims=True), 1e-9)


# ---------------------------------------------------------------- discovery
def course_of_src(p: Path):
    rel = p.relative_to(SRC)
    return CFG["courses"].get(rel.parts[0]) if len(rel.parts) > 1 else None


def refresh_courses():
    """auto_courses: every top-level folder of the source tree is a course (used for vaults whose modules are not listed in config)."""
    if not CFG.get("auto_courses"):
        return
    for d in sorted(SRC.iterdir()):
        if not d.is_dir() or d.name.startswith((".", "_")) or d.name in CFG["exclude_src"] or d.name in CFG["courses"]:
            continue
        slug = re.sub(r"[^A-Za-z0-9]+", "-", d.name).strip("-")
        if slug:
            CFG["courses"][d.name] = [slug, "00-" + slug, slug.lower()]


def wanted(p: Path):
    """src_all mode takes every file except the skip_exts list (video, audio...); otherwise only src_exts."""
    ext = p.suffix.lower()
    if p.name.startswith("~$"):
        return False
    if CFG.get("src_all"):
        return ext not in set(CFG.get("skip_exts", []))
    return ext in CFG["src_exts"]


def discover():
    """Yield (path, kind, course_tuple). Course tuple = (folder, hub, tag)."""
    refresh_courses()
    ex = [SRC / e for e in CFG["exclude_src"]]
    for p in sorted(SRC.rglob("*")):
        if not p.is_file() or not wanted(p):
            continue
        if any(e == p or e in p.parents for e in ex) or "__pycache__" in p.parts or p.name.startswith("."):
            continue
        if ".sync-conflict-" in p.name:
            continue
        c = course_of_src(p)
        if c:
            yield p, "doc", c
    by_folder = {v[0]: v for v in CFG["courses"].values()}
    for p in sorted(VAULT.rglob("*.md")):
        rel = p.relative_to(VAULT)
        if rel.parts[0] == "_auto" or len(rel.parts) < 2 or p.name in CFG["exclude_vault_names"]:
            continue
        if any(p.name.startswith(x) for x in CFG["exclude_vault_prefixes"]) or ".sync-conflict-" in p.name:
            continue
        c = by_folder.get(rel.parts[0])
        if c:
            yield p, "vault", c


# ---------------------------------------------------------------- extraction + chunking
class Skip(Exception):
    """File that cannot or should not be read: args[0] is the status to record (unsupported, too_big)."""


IMG_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
OFFICE_EXTS = {".pdf", ".docx", ".pptx", ".md"}
LANG = {".py": "python", ".r": "r", ".rmd": "r", ".sql": "sql", ".js": "javascript", ".ts": "typescript", ".java": "java",
        ".c": "c", ".h": "c", ".cpp": "cpp", ".hpp": "cpp", ".cs": "csharp", ".go": "go", ".rs": "rust", ".sh": "bash",
        ".json": "json", ".jsonl": "json", ".yaml": "yaml", ".yml": "yaml", ".toml": "toml", ".xml": "xml", ".html": "html",
        ".htm": "html", ".css": "css", ".tex": "latex", ".m": "matlab", ".jl": "julia", ".scala": "scala", ".kt": "kotlin",
        ".swift": "swift", ".ini": "ini", ".cfg": "ini", ".bib": "bibtex", ".csv": "csv", ".tsv": "tsv", ".txt": "text",
        ".log": "text", ".rst": "rst", ".ipynb": "python"}
TEXT_CAP = 400_000     # characters read from any one file
HEAD_ROWS = 60         # rows shown from a data table


def fenced(text, lang=""):
    fence = "`" * (max([len(m) for m in re.findall(r"`+", text)] + [2]) + 1)
    return f"{fence}{lang}\n{text.rstrip()}\n{fence}"


def read_plain(p: Path):
    raw = p.read_bytes()[: TEXT_CAP * 4]
    if b"\x00" in raw[:8192]:
        raise Skip("unsupported")   # binary
    text = raw.decode("utf-8", errors="replace")
    return text[:TEXT_CAP] + ("\n\n[truncated: file is longer]" if len(text) > TEXT_CAP else "")


def table_summary(p: Path, delim):
    """Header, first rows and row count of a delimited file, read as a stream so large datasets stay cheap."""
    import csv
    rows, n = [], 0
    with p.open(errors="replace", newline="") as f:
        for row in csv.reader(f, delimiter=delim):
            if n < HEAD_ROWS:
                rows.append(delim.join(row)[:400])
            n += 1
    return f"{n:,} rows (including the header). First {min(n, HEAD_ROWS)} shown.\n\n" + fenced("\n".join(rows), "csv")


def notebook_units(p: Path):
    nb = json.loads(p.read_text(errors="replace"))
    units = []
    for i, cell in enumerate(nb.get("cells", []), 1):
        src = "".join(cell.get("source", [])).strip()
        if not src:
            continue
        if cell.get("cell_type") == "markdown":
            units.append((f"cell {i}", src))
        elif cell.get("cell_type") == "code":
            out = ""
            for o in cell.get("outputs", []):
                t = "".join(o.get("text", [])) or "".join(o.get("data", {}).get("text/plain", []))
                if t.strip():
                    out += t.strip()[:1200] + "\n"
            units.append((f"cell {i}", fenced(src, "python") + (("\n\nOutput:\n" + fenced(out[:1500])) if out else "")))
    return units


def sheet_text(p: Path):
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets[:8]:
        rows = []
        for j, row in enumerate(ws.iter_rows(values_only=True)):
            if j >= HEAD_ROWS:
                break
            rows.append("\t".join("" if c is None else str(c) for c in row)[:400])
        parts.append(f"## Sheet {ws.title}\n\n" + fenced("\n".join(rows)))
    return "\n\n".join(parts)


def describe_image(p: Path):
    """Description of a standalone image by the vision model; cached by file hash so the raw export does not ask twice."""
    sha = sha_of(p)
    c = db()
    c.execute("CREATE TABLE IF NOT EXISTS image_text(sha TEXT PRIMARY KEY, text TEXT)")
    r = c.execute("SELECT text FROM image_text WHERE sha=?", (sha,)).fetchone()
    if r:
        return r[0]
    import base64
    unload(CFG["chat_model"])
    res = chat([{"role": "system", "content": "Describe this study image for a student's notes: what it shows, any text in it (copy it exactly), "
                 "axes, labels and the key technical idea. Return JSON."},
                {"role": "user", "content": "File: " + p.name}], VSCHEMA, model=CFG["vision_model"],
               images=[base64.b64encode(p.read_bytes()).decode()])
    text = f"Image ({res['visual_kind']}): {res['description']}\nConcepts: " + ", ".join(x["name"] for x in res["concepts"])
    c.execute("INSERT OR REPLACE INTO image_text VALUES(?,?)", (sha, text))
    c.commit()
    return text


def other_text(p: Path):
    """(extractor, markdown text, units) for every file that is not pdf, docx, pptx or md."""
    ext = p.suffix.lower()
    if p.stat().st_size > CFG.get("max_file_mb", 25) * 1_000_000:
        raise Skip("too_big")
    if ext in IMG_EXTS:
        t = describe_image(p)
        return "vision model", t, [("image", t)]
    if ext == ".ipynb":
        units = notebook_units(p)
        return "notebook cells", "\n\n".join(f"### {l}\n\n{t}" for l, t in units), units
    if ext in (".csv", ".tsv"):
        t = table_summary(p, "," if ext == ".csv" else "\t")
        return "table head", t, [("table", t)]
    if ext in (".xlsx", ".xlsm"):
        t = sheet_text(p)
        return "openpyxl", t, [("sheets", t)]
    if ext == ".zip":
        from markitdown import MarkItDown
        t = MarkItDown().convert(str(p)).text_content[:TEXT_CAP]
        if not t.strip():
            raise Skip("unsupported")
        return "markitdown (zip)", t, [("archive", t)]
    text = read_plain(p)
    lang = LANG.get(ext, "")
    body = fenced(text, lang) if lang not in ("", "text", "rst", "latex", "bibtex") or ext in (".log",) else text
    pieces, n = [], max(2000, CFG["chunk_chars"] - 500)
    for i in range(0, len(text), n):
        pieces.append((f"part {i // n + 1}", text[i:i + n]))
    return "plain text", body, pieces or [("empty", "")]


def pdf_pages(p: Path):
    """Page texts via pdfminer; a truncated or damaged PDF is rewritten by Ghostscript first and read again."""
    from pdfminer.high_level import extract_text
    try:
        return extract_text(str(p)).split("\f")
    except Exception:
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            fixed = Path(d) / "fixed.pdf"
            subprocess.run(["gs", "-q", "-o", str(fixed), "-sDEVICE=pdfwrite", str(p)], capture_output=True, timeout=300)
            if not fixed.exists() or fixed.stat().st_size < 1000:
                raise
            return extract_text(str(fixed)).split("\f")


def extract(p: Path):
    """Return list of (label, text) units: pages, slides or heading sections."""
    ext = p.suffix.lower()
    if ext not in OFFICE_EXTS:
        return other_text(p)[2]
    if ext == ".pdf":
        pages = pdf_pages(p)
        return [(f"p. {i + 1}", t) for i, t in enumerate(pages) if t.strip()]
    if ext == ".md":
        text = p.read_text(errors="ignore")
        text = re.sub(r"\A---\n.*?\n---\n", "", text, flags=re.S)
    else:
        from markitdown import MarkItDown
        text = MarkItDown().convert(str(p)).text_content
    if ext == ".pptx":
        parts = re.split(r"<!-- Slide number: (\d+) -->", text)
        return [(f"slide {parts[i]}", parts[i + 1]) for i in range(1, len(parts) - 1, 2) if parts[i + 1].strip()]
    secs, head, buf = [], "intro", []
    for line in text.splitlines():
        m = re.match(r"#{1,3}\s+(.*)", line)
        if m and buf and sum(map(len, buf)) > 400:
            secs.append((head, "\n".join(buf)))
            buf = []
        if m:
            head = m.group(1).strip()
        buf.append(line)
    if buf:
        secs.append((head, "\n".join(buf)))
    return secs


def chunk(units):
    out, label, buf = [], None, ""
    n = CFG["chunk_chars"]
    first = last = None
    for lab, t in units:
        t = re.sub(r"\n{3,}", "\n\n", t).strip()
        while len(t) > n * 1.5:  # split oversized unit
            cut = t.rfind("\n", 0, n)
            cut = cut if cut > n // 2 else n
            out.append((lab, t[:cut]))
            t = t[cut:].strip()
        if buf and len(buf) + len(t) > n:
            out.append((f"{first} to {last}" if first != last else first, buf))
            buf, first = "", None
        buf = (buf + "\n\n" + t).strip()
        first = first or lab
        last = lab
    if buf.strip():
        out.append((f"{first} to {last}" if first != last else first, buf))
    return [(h, t) for h, t in out if len(t) >= 200]


# ---------------------------------------------------------------- concept handling
def ckey(name):
    w = re.sub(r"[^a-z0-9]+", " ", name.lower()).split()
    if w and len(w[-1]) > 3 and w[-1].endswith("s") and not w[-1].endswith(("ss", "is", "us")):
        w[-1] = w[-1][:-1]
    return " ".join(w)


def norm_tokens(s):
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).split()


def quote_ok(quote, chunk_text):
    q = norm_tokens(quote)
    if len(q) < 3:
        return False
    ctoks = norm_tokens(chunk_text)
    frac = sum(t in set(ctoks) for t in q) / len(q)
    return frac >= 0.95 or (frac >= 0.85 and " ".join(q) in " ".join(ctoks))


SCHEMA = {"type": "object", "required": ["summary", "concepts"], "properties": {
    "summary": {"type": "string"},
    "concepts": {"type": "array", "items": {"type": "object",
        "required": ["name", "kind", "definition", "quote"], "properties": {
            "name": {"type": "string"},
            "kind": {"type": "string", "enum": ["concept", "method", "algorithm", "tool", "metric", "standard", "person", "model"]},
            "aliases": {"type": "array", "items": {"type": "string"}},
            "definition": {"type": "string"},
            "quote": {"type": "string"}}}}}}


def prompt(title, course, heading, text, known, hint=""):
    return [{"role": "system", "content":
        "You build a university study knowledge graph. From the text chunk, return JSON with a 1-3 sentence factual summary "
        "(British English, no fluff) and up to %d key technical concepts. Rules: name = the standard textbook name in singular form, "
        "full name first (put abbreviations such as SIFT or KF in aliases); the definition is at most 200 characters and must be supported "
        "by the text; quote = a verbatim sentence or phrase copied from the text that supports the concept. Skip admin content "
        "(deadlines, lecturer names, module logistics, slide numbers, generic words like figure, example, slide). "
        "If a concept matches one of the KNOWN names, use that exact name.\n" % RG["max_concepts_per_chunk"]
        + RULES.get("instructions", "")},
        {"role": "user", "content": "%sDocument: %s (course %s), section %s\nKNOWN: %s\n\nTEXT:\n%s" % (
            hint, title, course, heading, "; ".join(known) or "none", text)}]


def known_names(c, vec, k=25):
    rows = c.execute("SELECT name, emb FROM concepts WHERE emb IS NOT NULL").fetchall()
    if not rows:
        return []
    m = np.stack([np.frombuffer(r[1], dtype=np.float32) for r in rows])
    idx = np.argsort(-(m @ vec))[:k]
    return [rows[i][0] for i in idx]


IGNORE_KEYS = {ckey(x) for x in RULES.get("ignore", [])}
ALIAS_KEYS = {ckey(v): canon for canon, vs in (RULES.get("aliases") or {}).items() for v in vs}


def find_or_add_concept(c, name, kind, aliases):
    name = BAD_CHARS.sub("", name).strip()
    if not (2 <= len(name) <= 70):
        return None
    k = ckey(name)
    if not k or k in IGNORE_KEYS:
        return None
    if k in ALIAS_KEYS:
        name = ALIAS_KEYS[k]
        k = ckey(name)
    r = c.execute("SELECT id FROM concepts WHERE key=?", (k,)).fetchone()
    if r:
        return r[0]
    for a in aliases:  # alias match
        r = c.execute("SELECT id FROM concepts WHERE key=?", (ckey(a),)).fetchone()
        if r:
            return r[0]
    v = embed([name])[0]
    rows = c.execute("SELECT id, emb, kind FROM concepts WHERE emb IS NOT NULL").fetchall()
    if rows:
        m = np.stack([np.frombuffer(x[1], dtype=np.float32) for x in rows])
        sims = m @ v
        i = int(np.argmax(sims))
        if sims[i] >= CFG["merge_cosine"] and rows[i][2] == kind:
            al = json.loads(c.execute("SELECT aliases FROM concepts WHERE id=?", (rows[i][0],)).fetchone()[0])
            if name not in al:
                c.execute("UPDATE concepts SET aliases=? WHERE id=?", (json.dumps(al + [name]), rows[i][0]))
            return rows[i][0]
    cur = c.execute("INSERT INTO concepts(key,name,kind,aliases,emb) VALUES(?,?,?,?,?)",
                    (k, name, kind, json.dumps([a for a in aliases if a and ckey(a) != k][:5]), v.tobytes()))
    return cur.lastrowid


BOILER = re.compile(r"\b(thank you|any questions|questions\?|outline|agenda|learning outcomes|references|bibliography|"
                    r"contents|acknowledg)", re.I)


def is_boilerplate(text):
    return len(text) < 500 and bool(BOILER.search(text))


def cache_lookup(c, h, vec):
    r = c.execute("SELECT summary, items FROM chunk_cache WHERE hash=?", (h,)).fetchone()
    if not r:
        rows = c.execute("SELECT emb, summary, items FROM chunk_cache").fetchall()
        if rows:
            m = np.stack([np.frombuffer(x[0], dtype=np.float32) for x in rows])
            sims = m @ vec
            i = int(np.argmax(sims))
            if sims[i] >= RD["near_dup_cosine"]:
                r = rows[i][1:]
    return (r[0], [tuple(x) for x in json.loads(r[1])]) if r else None


def type_hint(path):
    ext = Path(path).suffix.lower()
    if ext in OFFICE_EXTS:
        return ""
    if ext in IMG_EXTS:
        return "FILE TYPE: a standalone image described by a vision model. Extract only the technical concepts it shows.\n"
    if ext in (".csv", ".tsv", ".xlsx", ".xlsm"):
        return ("FILE TYPE: a data table (header and first rows only). Concepts = the kind of data, the variables and the analysis it "
                "suggests; do not list individual values.\n")
    return ("FILE TYPE: source code, a notebook or a data file. Concepts = the algorithms, statistical methods, libraries and techniques it "
            "demonstrates, named as a textbook would (for example 'k-means clustering', not 'KMeans(n_clusters=3)').\n")


def process_doc(c, doc_id, path, kind, course):
    title = c.execute("SELECT title FROM docs WHERE id=?", (doc_id,)).fetchone()[0]
    try:
        chunks = chunk(extract(path))
    except Skip as e:
        c.execute("UPDATE docs SET status=? WHERE id=?", (str(e), doc_id))
        c.commit()
        log(f"skipped {path.name}: {e}")
        return
    if not chunks:
        c.execute("UPDATE docs SET status='no_text' WHERE id=?", (doc_id,))
        c.commit()
        return
    done = {r[0] for r in c.execute("SELECT idx FROM chunks WHERE doc_id=?", (doc_id,))}
    for i, (heading, text) in enumerate(chunks):
        if i in done:
            continue
        if RD["skip_boilerplate"] and is_boilerplate(text):
            c.execute("INSERT OR REPLACE INTO chunks VALUES(?,?,?,?)", (doc_id, i, heading, ""))
            c.commit()
            continue
        vec = embed([text[:1500]])[0]
        h = hashlib.sha1(" ".join(norm_tokens(text)).encode()).hexdigest()
        hit = cache_lookup(c, h, vec)
        if hit:
            summary, items = hit
            for cid, quote, defn in items:
                c.execute("INSERT OR REPLACE INTO mentions VALUES(?,?,?,?,?)", (cid, doc_id, i, quote, defn))
            c.execute("INSERT OR REPLACE INTO chunks VALUES(?,?,?,?)", (doc_id, i, heading, summary))
            c.commit()
            log(f"  {title[:50]} chunk {i + 1}/{len(chunks)}: reused (duplicate content, no model call)")
            continue
        res = chat(prompt(title, course[0], heading, text, known_names(c, vec), type_hint(path)), SCHEMA)
        if res is None:
            raise Offline("model returned no valid JSON after retries")
        n, items = 0, []
        for con in res.get("concepts", [])[:RG["max_concepts_per_chunk"]]:
            if not quote_ok(con.get("quote", ""), text):
                continue
            cid = find_or_add_concept(c, con["name"], con["kind"], con.get("aliases", []))
            if cid:
                items.append((cid, con["quote"][:240], con["definition"][:260]))
                c.execute("INSERT OR REPLACE INTO mentions VALUES(?,?,?,?,?)", (cid, doc_id, i) + items[-1][1:])
                n += 1
        c.execute("INSERT OR REPLACE INTO chunks VALUES(?,?,?,?)", (doc_id, i, heading, res.get("summary", "")))
        c.execute("INSERT OR REPLACE INTO chunk_cache VALUES(?,?,?,?)", (h, vec.tobytes(), res.get("summary", ""), json.dumps(items)))
        c.commit()
        log(f"  {title[:50]} chunk {i + 1}/{len(chunks)}: {n} concepts")
        breathe(RT["sleep_between_calls_s"])
    c.execute("UPDATE docs SET status='done' WHERE id=?", (doc_id,))
    c.commit()


# ---------------------------------------------------------------- raw text export (no AI)
RAWDIR = VAULT / "_raw"


def raw_text(p: Path):
    """Full text exactly as the extractor returns it, with page or slide headings. No model involved."""
    ext = p.suffix.lower()
    if ext == ".pdf":
        pages = pdf_pages(p)
        if pages and not pages[-1].strip():
            pages.pop()
        return "pdfminer.six", "\n\n".join(f"## Page {i + 1}\n\n{t.strip() or '(no text layer on this page)'}"
                                          for i, t in enumerate(pages))
    if ext == ".md":
        text = p.read_text(errors="ignore")
        # verbatim, but inside a code fence so the original [[wikilinks]] do not become links in your graph
        fence = "`" * (max([len(m) for m in re.findall(r"`+", text)] + [2]) + 1)
        return "verbatim copy (in a code fence)", f"{fence}markdown\n{text.rstrip()}\n{fence}"
    if ext not in OFFICE_EXTS:
        extractor, text, _ = other_text(p)
        return extractor, text
    from markitdown import MarkItDown
    text = MarkItDown().convert(str(p)).text_content
    if ext == ".pptx":
        text = re.sub(r"<!-- Slide number: (\d+) -->", r"## Slide \1", text)
    return "markitdown", text


def raw_export(c, vision_ok=True):
    """One untouched full-text note per source file under _raw/<course>/raw-<name>.md."""
    n = 0
    for did, path, course, title, sha, doctype in c.execute(
            "SELECT id,path,course,title,sha,doctype FROM docs WHERE kind='doc' AND status!='excluded'").fetchall():
        r = c.execute("SELECT sha FROM raw_exports WHERE doc_id=?", (did,)).fetchone()
        if r and r[0] == sha:
            continue
        if Path(path).suffix.lower() in IMG_EXTS and not vision_ok:
            continue   # describing an image needs the vision model, which only the vault holding the model lock may load
        course = json.loads(course)
        try:
            extractor, text = raw_text(Path(path))
        except Skip:
            continue   # unreadable or too large: nothing to export
        except Exception as e:
            log("raw export failed:", Path(path).name, e)
            continue
        name = BAD_CHARS.sub("", title).strip()[:90]
        rel = f"_raw/{course[0]}/raw-{name}.md"
        if c.execute("SELECT 1 FROM raw_exports WHERE path=? AND doc_id!=?", (rel, did)).fetchone():
            rel = f"_raw/{course[0]}/raw-{name} [{doctype}].md"
        slug = BAD_CHARS.sub("", title).strip()[:80] + " " + sha[:4]
        src_rel = Path(path).relative_to(SRC)
        body = "\n".join([
            "---", "type: raw", f"doctype: {doctype}", f"course: {course[2]}", f"tags: [raw, {doctype}, {course[2]}]",
            f"source_file: {json.dumps(str(src_rel))}", f"sha256: {sha}", f"extractor: {extractor}",
            f"extracted: {time.strftime('%Y-%m-%d')}", "---", f"# raw-{name}", "",
            "Unprocessed text extraction. No AI has read, summarised or changed this file. Images, diagrams and "
            "equations drawn as pictures are not included.  ",
            f"Source file: `{src_rel}`  ", f"Processed note: [[{VP}_auto/sources/{course[0]}/{slug}|{title}]]", "",
            "---", "", text.rstrip(), ""])
        out = VAULT / rel
        out.parent.mkdir(parents=True, exist_ok=True)
        tmp = out.with_name("." + out.name + ".tmp")
        tmp.write_text(body)
        tmp.replace(out)
        c.execute("INSERT OR REPLACE INTO raw_exports VALUES(?,?,?)", (did, sha, rel))
        c.commit()
        n += 1
    if n:
        log(f"raw export: wrote {n} untouched text notes")


# ---------------------------------------------------------------- vision pass
VSCHEMA = {"type": "object", "required": ["visual_kind", "description", "concepts"], "properties": {
    "visual_kind": {"type": "string", "enum": ["diagram", "plot", "equation", "table", "photo", "screenshot",
                                                "text_slide"]},
    "description": {"type": "string"},
    "concepts": SCHEMA["properties"]["concepts"]}}


def figure_candidates(path: Path, units):
    """Pages worth sending to the vision model: a large non-repeating image, or almost no text."""
    import subprocess as sp
    info = sp.run(["pdfinfo", str(path)], capture_output=True, text=True).stdout
    m = re.search(r"Pages:\s+(\d+)", info)
    total = int(m.group(1)) if m else 0
    objpages, rows = Counter(), []
    for ln in sp.run(["pdfimages", "-list", str(path)], capture_output=True, text=True).stdout.splitlines()[2:]:
        f = ln.split()
        if len(f) > 10 and f[2] == "image" and int(f[3]) >= 200 and int(f[4]) >= 200:
            rows.append((int(f[0]), f[10]))
            objpages[f[10]] += 1
    withimg = {pg for pg, obj in rows if objpages[obj] <= 2}  # logos and template art repeat on many pages
    textlen = {int(lab.split()[1]): len(t) for lab, t in units if lab.startswith("p. ")}
    return sorted(pg for pg in range(1, total + 1) if pg in withimg or textlen.get(pg, 0) < 250)


def plan_figures(c):
    for did, path in c.execute("SELECT id, path FROM docs WHERE status='done' AND doctype='pdf' AND kind='doc' "
                               "AND id NOT IN (SELECT DISTINCT doc_id FROM figure_pages)").fetchall():
        try:
            pages = figure_candidates(Path(path), extract(Path(path)))
        except Exception as e:
            log("figure planning failed", Path(path).name, e)
            pages = []
        c.executemany("INSERT OR IGNORE INTO figure_pages VALUES(?,?, 'todo')", [(did, pg) for pg in pages])
        c.execute("INSERT OR IGNORE INTO figure_pages VALUES(?, 0, 'planned')", (did,))  # marker row
    c.commit()


def lookup_concept(c, name, thr=0.85):
    k = ckey(name)
    r = c.execute("SELECT id FROM concepts WHERE key=?", (k,)).fetchone()
    if r:
        return r[0]
    rows = c.execute("SELECT id, emb FROM concepts WHERE emb IS NOT NULL").fetchall()
    if not rows:
        return None
    m = np.stack([np.frombuffer(x[1], dtype=np.float32) for x in rows])
    sims = m @ embed([name])[0]
    i = int(np.argmax(sims))
    return rows[i][0] if sims[i] >= thr else None


def vision_phase(c, limit=0):
    import base64, tempfile
    plan_figures(c)
    todo = c.execute("SELECT f.doc_id, f.page, d.path, d.title, d.course FROM figure_pages f JOIN docs d ON d.id=f.doc_id "
                     "WHERE f.status='todo' ORDER BY f.doc_id, f.page").fetchall()
    if not todo:
        return
    log(f"vision pass: {len(todo)} pages")
    unload(CFG["chat_model"])
    figdir = AUTO / "figures"
    figdir.mkdir(parents=True, exist_ok=True)
    try:
        for n, (did, page, path, title, course) in enumerate(todo, 1):
            if limit and n > limit:
                break
            course = json.loads(course)
            with tempfile.TemporaryDirectory() as td:
                subprocess.run(["pdftoppm", "-r", str(RT["vision_dpi"]), "-f", str(page), "-l", str(page), "-jpeg",
                                "-jpegopt", "quality=80", "-singlefile", path, f"{td}/pg"], check=True)
                raw = Path(f"{td}/pg.jpg").read_bytes()
            ptext = ""
            for lab, t in extract(Path(path)):
                if lab == f"p. {page}":
                    ptext = t
            vec = embed([ptext[:1500] or title])[0]
            msgs = [{"role": "system", "content":
                     "You read one lecture slide image for a university study knowledge graph. Return JSON. visual_kind: what the "
                     "slide mainly shows; any plot, graph, curve, diagram, equation or table counts as such even if the slide also has "
                     "text; use text_slide only for slides that are just words, a title card or a logo. description: 2-4 "
                     "sentences in British English saying what it teaches, naming labelled parts, axes and trends, with equations "
                     "in LaTeX; do not invent numbers you cannot read. concepts: at most %d key technical concepts shown "
                     "(name, kind, aliases, short definition, quote = a label or phrase visible on the slide). If a concept matches "
                     "a KNOWN name use that exact name.\n%s" % (min(RG["max_concepts_per_chunk"], 4), RULES.get("instructions", ""))},
                    {"role": "user", "content": "Document: %s (course %s), page %d\nKNOWN: %s\nText layer of this page: %s" % (
                        title, course[0], page, "; ".join(known_names(c, vec)), ptext[:1200] or "none")}]
            res = chat(msgs, VSCHEMA, model=CFG["vision_model"], images=[base64.b64encode(raw).decode()])
            if res is None:
                raise Offline("vision model returned no valid JSON")
            idx = 1000 + page
            if res["visual_kind"] == "text_slide" and len(ptext) >= 250:  # words only; the text pass already has it
                res["concepts"] = []
            if res["visual_kind"] == "text_slide" and not res["concepts"]:
                c.execute("UPDATE figure_pages SET status='skipped' WHERE doc_id=? AND page=?", (did, page))
                c.commit()
                log(f"  {title[:40]} p.{page}: words only, skipped")
            else:
                fname = f"{did:03d}-p{page}.jpg"
                tmpf = figdir / ("." + fname + ".tmp")
                tmpf.write_bytes(raw)
                tmpf.replace(figdir / fname)
                pool = set(norm_tokens(ptext))
                kept = 0
                for con in res["concepts"][:4]:
                    name_toks = norm_tokens(con["name"])
                    cid = None
                    if name_toks and all(t in pool for t in name_toks):
                        cid = find_or_add_concept(c, con["name"], con["kind"], con.get("aliases", []))
                    if cid is None:
                        cid = lookup_concept(c, con["name"])
                    if cid:
                        c.execute("INSERT OR REPLACE INTO mentions VALUES(?,?,?,?,?)",
                                  (cid, did, idx, "(from figure) " + con["quote"][:200], con["definition"][:260]))
                        kept += 1
                c.execute("INSERT OR REPLACE INTO chunks VALUES(?,?,?,?)",
                          (did, idx, f"p. {page} (figure: {res['visual_kind']}) [{fname}]", res["description"]))
                c.execute("UPDATE figure_pages SET status='done' WHERE doc_id=? AND page=?", (did, page))
                c.commit()
                log(f"  {title[:40]} p.{page}: {res['visual_kind']}, {kept} concepts")
            breathe(RT["vision_sleep_s"])
    finally:
        unload(CFG["vision_model"])


# ---------------------------------------------------------------- scan
def scan(c):
    if CFG.get("require_marker") and not (SRC / ".stfolder").exists():
        raise Offline("source folder marker .stfolder missing (require_marker is on); not scanning")
    seen = set()
    shas = {r[0] for r in c.execute("SELECT sha FROM docs WHERE kind='doc'")}
    for p, kind, course in discover():
        key = str(p)
        seen.add(key)  # counts as present even if we hold off processing it
        if time.time() - p.stat().st_mtime < 60:
            continue  # may still be arriving from sync; pick it up on a later run
        sha = sha_of(p)
        row = c.execute("SELECT id, sha FROM docs WHERE path=?", (key,)).fetchone()
        if row and row[1] == sha:
            continue
        if not row and kind == "doc" and sha in shas:
            continue  # duplicate copy of an ingested file
        if row:  # changed: drop its old chunks and mentions
            for t in ("chunks", "mentions"):
                c.execute(f"DELETE FROM {t} WHERE doc_id=?", (row[0],))
            c.execute("UPDATE docs SET sha=?, status=CASE WHEN status='excluded' THEN 'excluded' ELSE 'pending' END WHERE id=?", (sha, row[0]))
        else:
            title = p.stem.replace("_", " ")
            if kind == "vault":
                link = str(p.relative_to(VAULT).with_suffix(""))
            else:
                link = ""
            # standalone images need the vision model; queue them last so they run as one batch instead of swapping models per file
            c.execute("INSERT INTO docs(path,sha,kind,course,title,status,doctype,link,priority) VALUES(?,?,?,?,?,'pending',?,?,?)",
                      (key, sha, kind, json.dumps(course), title, p.suffix.lower().lstrip("."), link, -1 if p.suffix.lower() in IMG_EXTS else 0))
            shas.add(sha)
    # files removed from disk: forget them (rendered notes disappear on next render)
    gone = [(path, did) for (path, did) in c.execute("SELECT path, id FROM docs").fetchall() if path not in seen]
    total = c.execute("SELECT COUNT(*) FROM docs").fetchone()[0]
    if gone and len(gone) > max(3, 0.25 * total):
        log(f"WARNING: {len(gone)} of {total} known files vanished at once; treating as a sync problem, keeping state")
        gone = []
    for path, did in gone:  # files really removed: forget them (their generated notes disappear on next render)
        for t in ("chunks", "mentions"):
            c.execute(f"DELETE FROM {t} WHERE doc_id=?", (did,))
        c.execute("DELETE FROM docs WHERE id=?", (did,))
    c.commit()


# ---------------------------------------------------------------- render
QUOTE_MARKS = ' "\'\u201c\u201d\u2018\u2019'   # models often wrap the quote in its own marks


def wl(link, label=None):
    return f"[[{link}|{label}]]" if label else f"[[{link}]]"


def write_managed(path: Path, generated: str, written: set):
    """Write generated text, preserving anything the user added below USER_MARK."""
    tail = ""
    if path.exists():
        old = path.read_text()
        if USER_MARK in old:
            tail = old.split(USER_MARK, 1)[1]
    body = generated.rstrip() + "\n\n" + USER_MARK + (tail if tail.strip() else "\n")
    written.add(path)
    if not path.exists() or path.read_text() != body:
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_name("." + path.name + ".tmp")  # dot-prefixed: sync and Obsidian ignore it
        tmp.write_text(body)
        tmp.replace(path)


def render(c):
    mn = CFG["min_mentions_for_note"]
    docs = {r[0]: r for r in c.execute("SELECT id,path,kind,course,title,doctype,link,sha FROM docs")}
    cons = {r[0]: r for r in c.execute("SELECT id,name,kind,aliases FROM concepts")}
    ment = defaultdict(list)
    for r in c.execute("SELECT concept_id,doc_id,idx,quote,definition FROM mentions"):
        if r[1] in docs:
            ment[r[0]].append(r)
    def keep_ok(cid, ms):
        nd = len({m[1] for m in ms})
        if cons[cid][2] == "person" and nd < 3:
            return False
        return nd >= RG["min_docs_for_note"] or len(ms) >= RG["min_chunks_alt"]

    keep = {cid for cid, ms in ment.items() if keep_ok(cid, ms)}
    names = {cid: BAD_CHARS.sub("", cons[cid][1]) for cid in keep}
    chunks = {(r[0], r[1]): r for r in c.execute("SELECT doc_id,idx,heading,summary FROM chunks")}
    rawpath = dict(c.execute("SELECT doc_id, path FROM raw_exports"))

    def target(doc):  # wikilink target and label for a doc
        course = json.loads(doc[3])
        if doc[2] == "vault":
            return VP + doc[6], doc[4]
        return f"{VP}_auto/sources/{course[0]}/{src_slug(doc)}", doc[4]

    def src_slug(doc):
        return BAD_CHARS.sub("", doc[4]).strip()[:80] + " " + doc[7][:4]

    cooc = defaultdict(Counter)
    by_chunk = defaultdict(list)
    for cid in keep:
        for m in ment[cid]:
            by_chunk[(m[1], m[2])].append(cid)
    for cs in by_chunk.values():
        for a in cs:
            for b in cs:
                if a != b:
                    cooc[a][b] += 1

    written = set()
    course_concepts = defaultdict(Counter)
    concept_courses = defaultdict(set)
    for cid in keep:
        for m in ment[cid]:
            course = json.loads(docs[m[1]][3])
            concept_courses[cid].add(course[2])
            course_concepts[course[2]][cid] += 1
    tag_to_course = {v[2]: v for v in CFG["courses"].values()}

    # concept notes
    for cid in keep:
        cid_, name, kind, aliases = cons[cid]
        ms = ment[cid]
        best = max(ms, key=lambda m: len(m[4]))
        tags = list(dict.fromkeys(["concept", kind] + sorted(concept_courses[cid]) + (["cross-course"] if len(concept_courses[cid]) >= 2 else [])))
        fm = ["---", "type: concept", f"kind: {kind}", f"mentions: {len(ms)}",
              "courses: [" + ", ".join(sorted(concept_courses[cid])) + "]",
              "tags: [" + ", ".join(tags) + "]"]
        al = [a for a in json.loads(aliases) if a]
        if al:
            fm.append("aliases: [" + ", ".join(json.dumps(a) for a in al) + "]")
        fm.append("---")
        lines = fm + [f"# {names[cid]}", "", f"> {best[4]}", "", "## Appears in"]
        seen_t = set()
        for m in sorted(ms, key=lambda m: (docs[m[1]][2] != "vault", docs[m[1]][4])):
            t, lab = target(docs[m[1]])
            head = chunks.get((m[1], m[2]), (0, 0, ""))[2]
            if (t, head) in seen_t:
                continue
            seen_t.add((t, head))
            lines.append(f"- {wl(t, lab)} ({head}): \"{m[3].strip(QUOTE_MARKS)}\"")
            if len(seen_t) >= RG["max_sources_per_concept"]:
                break
        rel = [b for b, n in cooc[cid].most_common(RG["related_per_concept"] * 3)
               if b in keep and n >= RG["related_min_cooccur"]][:RG["related_per_concept"]]
        if rel:
            lines += ["", "## Related concepts"] + [f"- {wl(names[b])}" for b in rel]
        write_managed(AUTO / "concepts" / f"{names[cid]}.md", "\n".join(lines), written)

    # source notes (external docs only; vault notes stay untouched)
    for did, doc in docs.items():
        if doc[2] != "doc":
            continue
        course = json.loads(doc[3])
        dch = sorted(k for k in chunks if k[0] == did)
        if not dch:
            continue
        mine = defaultdict(list)
        for cid in keep:
            for m in ment[cid]:
                if m[1] == did:
                    mine[m[2]].append(cid)
        used = Counter(cid for v in mine.values() for cid in v)
        linked = {cid for cid, _ in sorted(used.items(), key=lambda x: -len(ment[x[0]]))[:RG["max_concept_links_per_source"]]}
        related = Counter()
        for cid in used:
            for m in ment[cid]:
                if docs[m[1]][2] == "vault":
                    related[m[1]] += 1
        rel_path = Path(doc[1]).relative_to(SRC)
        lines = ["---", "type: source", f"doctype: {doc[5]}", f"course: {course[2]}",
                 "tags: [source, " + doc[5] + ", " + course[2] + "]", f"source_file: {json.dumps(str(rel_path))}", "---",
                 f"# {doc[4]}", "", f"Course hub: {wl(VP + course[0] + '/' + course[1], course[1])}  ",
                 f"Source file: `{rel_path}`"]
        if did in rawpath:
            lines.append(f"Raw text (no AI): {wl(VP + rawpath[did][:-3], 'raw-' + doc[4])}")
        lines += ["", "## Sections"]
        for k in dch:
            ch = chunks[k]
            fig = re.search(r"\[(\d+-p\d+\.jpg)\]$", ch[2])
            if fig:
                lines += [f"### {ch[2][:fig.start()].strip()}", f"![[{VP}_auto/figures/{fig.group(1)}|520]]", "", ch[3] or "",
                          "_Machine-read from the slide image; check equations and numbers against the source._", ""]
            else:
                lines += [f"### {ch[2]}", ch[3] or "", ""]
            if mine.get(k[1]):
                lines.append("Concepts: " + ", ".join(wl(names[cid]) if cid in linked else names[cid]
                                                    for cid in dict.fromkeys(mine[k[1]])))
                lines.append("")
        top = [t for t, n in related.most_common(8) if n >= 2]
        if top:
            lines += ["## Related topic notes"] + [f"- {wl(VP + docs[t][6], docs[t][4])}" for t in top]
        write_managed(AUTO / "sources" / course[0] / (src_slug(doc) + ".md"), "\n".join(lines), written)

    # per-course concept maps + top-level map
    for tag, cc in course_concepts.items():
        folder, hub, _ = tag_to_course[tag]
        lines = ["---", "type: map", f"tags: [map, {tag}]", "---", f"# {folder} concept map", "",
                 f"Hub: {wl(VP + folder + '/' + hub, hub)}", "", "## Concepts by mentions"]
        lines += [f"- {wl(names[cid])} ({n})" for cid, n in cc.most_common(80)]
        write_managed(AUTO / "maps" / f"{folder} concept map.md", "\n".join(lines), written)
    bridges = sorted((cid for cid in keep if len(concept_courses[cid]) >= 2),
                     key=lambda x: (-len(concept_courses[x]), -len(ment[x])))
    lines = ["---", "type: map", "tags: [map, knowledge-graph]", "---", "# Knowledge graph", "",
             "Generated by vault-agent. Concepts below appear in more than one course and bridge the graph.", "",
             "## Courses"]
    lines += [f"- {wl(VP + '_auto/maps/' + v[0] + ' concept map', v[0])}" if v[2] in course_concepts else f"- {v[0]}"
              for v in CFG["courses"].values()]
    lines += ["", "## Cross-course concepts"]
    lines += [f"- {wl(names[cid])} ({', '.join(sorted(concept_courses[cid]))})" for cid in bridges[:100]]
    write_managed(AUTO / "Knowledge graph.md", "\n".join(lines), written)

    # a starter hub note per course, created once and then left to the user (generated notes link to it)
    for folder, hub, tag in CFG["courses"].values():
        hp = VAULT / folder / (hub + ".md")
        if not hp.exists():
            hp.parent.mkdir(parents=True, exist_ok=True)
            hp.write_text(f"---\ntags: [hub, {tag}]\n---\n# {folder}\n\nYour notes for this course go in this folder. "
                          f"The agent links here but never edits this note.\n\n"
                          f"- Concept map: {wl(VP + '_auto/maps/' + folder + ' concept map', 'concept map')}\n"
                          f"- Knowledge graph: {wl(VP + '_auto/Knowledge graph', 'all courses')}\n")

    # remove stale generated files (only those we wrote in earlier runs, never anything else)
    stale = [p for p in AUTO.rglob("*.md") if p not in written and USER_MARK in p.read_text()]
    existing = sum(1 for _ in AUTO.rglob("*.md"))
    if len(stale) > 10 and len(stale) > 0.3 * max(existing, 1):
        log(f"WARNING: render would delete {len(stale)} of {existing} generated notes; skipping deletion")
    else:
        for p in stale:
            p.unlink()
    for d in sorted((x for x in AUTO.rglob("*") if x.is_dir()), reverse=True):
        if not any(d.iterdir()):
            d.rmdir()
    return len(keep), len(bridges)


# ---------------------------------------------------------------- backup
def backup(msg):
    root = cfg_path(CFG["vault_root"])
    def git(*a):
        return subprocess.run(["git", "-C", str(root), *a], capture_output=True, text=True)
    if not (root / ".git").exists():
        git("init", "-q", "-b", "main")
        (root / ".gitignore").write_text(".obsidian/workspace*.json\n.obsidian/cache\n.trash/\n")
        if CFG.get("backup_remote"):
            git("remote", "add", "origin", CFG["backup_remote"])
    git("add", "-A")
    if git("status", "--porcelain").stdout.strip():
        git("-c", "user.name=vault-agent", "-c", "user.email=vault-agent@localhost", "commit", "-q", "-m", msg)
    ok = True
    if CFG.get("backup_remote"):
        ok = git("push", "-q", "origin", "main").returncode == 0
    if CFG.get("state_backup"):
        subprocess.run(["rsync", "-a", "--mkpath", str(STATE / "state.db"), CFG["state_backup"]], capture_output=True)
    if CFG.get("backup_dir"):
        bdir = cfg_path(CFG["backup_dir"])
        bdir.mkdir(parents=True, exist_ok=True)
        dest = bdir / f"state-{time.strftime('%Y%m%d')}.db"
        if not dest.exists():  # one consistent snapshot per day, last 7 kept
            src = sqlite3.connect(STATE / "state.db")
            dst = sqlite3.connect(dest)
            src.backup(dst)
            dst.close()
            src.close()
            for old in sorted(bdir.glob("state-*.db"))[:-7]:
                old.unlink()
    return ok


def write_status(c, phase):
    counts = dict(c.execute("SELECT status, COUNT(*) FROM docs GROUP BY status").fetchall())
    figs = dict(c.execute("SELECT status, COUNT(*) FROM figure_pages WHERE page > 0 GROUP BY status").fetchall())
    ncon = c.execute("SELECT COUNT(*) FROM concepts").fetchone()[0]
    nm = c.execute("SELECT COUNT(*) FROM mentions").fetchone()[0]
    lines = ["---", "type: status", "tags: [map, status]", "---", "# Vault agent status", "",
             f"- Last update: {time.strftime('%Y-%m-%d %H:%M')} (server clock)", f"- Phase: {phase}",
             f"- Documents: {counts}", f"- Figure pages: {figs}", f"- Concepts: {ncon}, mentions: {nm}",
             "- If the last update is more than a day old, the vault agent is not running.", ""]
    path = AUTO / "Status.md"
    path.parent.mkdir(parents=True, exist_ok=True)
    body = "\n".join(lines)
    if not path.exists() or path.read_text() != body:
        tmp = path.with_name(".Status.md.tmp")
        tmp.write_text(body)
        tmp.replace(path)


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["run", "render", "status"], default="run", nargs="?")
    ap.add_argument("--limit", type=int, default=0, help="max docs to process this run")
    a = ap.parse_args()
    SRC.mkdir(parents=True, exist_ok=True)
    VAULT.mkdir(parents=True, exist_ok=True)
    lock = open(STATE / "lock", "w")
    try:
        if a.cmd != "status":
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError:
        log("another run is active")
        return
    c = db()
    if a.cmd == "status":
        for st, n in c.execute("SELECT status, COUNT(*) FROM docs GROUP BY status"):
            print(st, n)
        print("concepts", c.execute("SELECT COUNT(*) FROM concepts").fetchone()[0],
              "mentions", c.execute("SELECT COUNT(*) FROM mentions").fetchone()[0])
        return
    scan(c)   # no model needed: files appear in the queue even while another vault is using the model
    have_llm = False
    if a.cmd == "run":
        have_llm = True
        if CFG.get("shared_llm_lock"):   # one big model at a time across all vaults on this machine
            biglock = open(cfg_path(CFG["shared_llm_lock"]), "w")
            try:
                fcntl.flock(biglock, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except BlockingIOError:
                have_llm = False
                log("another vault is using the model, so documents wait (new files are still listed)")
    if a.cmd in ("run", "render"):
        raw_export(c, vision_ok=have_llm)
    if a.cmd == "run" and have_llm:
        pending_n = c.execute("SELECT COUNT(*) FROM docs WHERE status='pending'").fetchone()[0]
        log(f"{pending_n} documents pending")
        if pending_n or c.execute("SELECT 1 FROM figure_pages WHERE status='todo'").fetchone() or \
                c.execute("SELECT 1 FROM docs WHERE status='done' AND doctype='pdf' AND kind='doc' AND id NOT IN "
                          "(SELECT DISTINCT doc_id FROM figure_pages)").fetchone():
            try:
                ensure_tunnel()
                need_models("chat_model", "embed_model")
                free_study_models(True)
                tried, n = {0}, 0
                while True:   # the next document is chosen fresh each time, so queue changes made in STELLA-80 apply at once
                    row = c.execute("SELECT id,path,kind,course FROM docs WHERE status='pending' AND id NOT IN (%s) "
                                    "ORDER BY (kind='vault') DESC, priority DESC, id LIMIT 1" % ",".join(map(str, tried))).fetchone()
                    if not row:
                        break
                    did, path, kind, course = row
                    n += 1
                    if a.limit and n > a.limit:
                        break
                    tried.add(did)
                    log(f"[{n}/{max(pending_n, n)}] {Path(path).name}")
                    try:
                        process_doc(c, did, Path(path), kind, json.loads(course))
                    except Offline:
                        raise
                    except Exception as e:   # one unreadable file (corrupt PDF, bad encoding) must not stall the whole queue
                        c.execute("UPDATE docs SET status='error' WHERE id=?", (did,))
                        c.commit()
                        log(f"failed {Path(path).name}: {type(e).__name__}: {str(e)[:120]}")
                    if n % 5 == 0:   # pick up files and folders that arrived while this long run was going
                        scan(c)
                        pending_n = max(pending_n, n + c.execute("SELECT COUNT(*) FROM docs WHERE status='pending'").fetchone()[0])
                    k, b = render(c)
                    write_status(c, f"text pass, document {n} of {max(pending_n, n)} this run")
                    if n % 5 == 0:
                        backup(f"vault-agent: {k} concepts")
                if not c.execute("SELECT 1 FROM docs WHERE status='pending'").fetchone():
                    ensure_tunnel()
                    need_models("vision_model", "embed_model")
                    vision_phase(c, a.limit)
            except Offline as e:
                log("stopping:", e, "(will retry on next timer run)")
    k, b = render(c)
    pend = c.execute("SELECT COUNT(*) FROM docs WHERE status='pending'").fetchone()[0]
    write_status(c, "text pass" if pend else "idle or vision pass")
    ok = backup(f"vault-agent: {k} concepts, {b} cross-course")
    log(f"rendered {k} concept notes, {b} cross-course; backup {'ok' if ok else 'FAILED'}")


if __name__ == "__main__":
    main()
